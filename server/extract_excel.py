import openpyxl, json, os

path = 'SERRANO ACTUALIZADO 15.xlsx' if os.path.exists('SERRANO ACTUALIZADO 15.xlsx') else 'server/data/temp.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)

# 1. Proyectos
ws_p = wb['PROYECTOS']
proyectos = []
for r in ws_p.iter_rows(min_row=2):
    if not r[0].value: continue
    proyectos.append({
        'codigo': str(r[0].value).strip().upper(),
        'nombre': str(r[1].value).strip(),
        'municipio': str(r[2].value).strip() if r[2].value else None,
        'lotes_planeados': int(r[3].value) if r[3].value else None
    })

# 2. Lotes y Ventas
ws_v = wb['VENTAS']
lotes = []
ventas = []
for r in ws_v.iter_rows(min_row=2):
    if not r[0].value: continue
    id_venta = str(r[0].value).strip()
    id_proy = str(r[1].value).strip() if r[1].value else 'Ama'
    lote = str(r[2].value).strip() if r[2].value else ''
    area = float(r[3].value) if r[3].value else None
    comprador = str(r[4].value).strip() if r[4].value else None
    cedula = str(r[5].value).strip() if r[5].value else None
    tel = str(r[6].value).strip() if r[6].value else None
    correo = str(r[7].value).strip() if r[7].value else None
    dir_cli = str(r[8].value).strip() if r[8].value else None
    fecha = str(r[9].value)[:10] if r[9].value else None
    valor = float(r[10].value) if r[10].value else None
    inicial = float(r[11].value) if r[11].value else None
    obs = str(r[14].value).strip() if r[14].value else None

    lotes.append({'id_proy': id_proy, 'codigo': lote, 'area_m2': area})
    if comprador and valor and valor > 0:
        ventas.append({
            'codigo_venta': id_venta,
            'id_proy': id_proy,
            'codigo_lote': lote,
            'comprador': comprador,
            'cedula': cedula or ('SIN-DOC-' + id_venta),
            'telefono': tel,
            'email': correo,
            'direccion': dir_cli,
            'fecha': fecha or '2026-08-01',
            'valor': valor,
            'cuota_inicial': inicial,
            'observaciones': obs
        })

# 3. Abonos
ws_a = wb['ABONOS']
abonos = []
for r in ws_a.iter_rows(min_row=2):
    if not r[0].value or not r[1].value: continue
    id_abono = str(r[0].value).strip()
    id_venta = str(r[1].value).strip()
    fecha = str(r[3].value)[:10] if r[3].value else None
    valor = float(r[4].value) if r[4].value else None
    medio_raw = str(r[5].value).strip() if r[5].value else ''
    ref = str(r[6].value).strip() if r[6].value else None
    obs = str(r[7].value).strip() if r[7].value else None

    # Normalizar medio
    medio = 'otro'
    m_low = medio_raw.lower()
    if 'transf' in m_low: medio = 'transferencia'
    elif 'consign' in m_low: medio = 'consignacion'
    elif 'efect' in m_low: medio = 'efectivo'
    elif not medio_raw: medio = None

    if valor and valor > 0:
        abonos.append({
            'id_abono': id_abono,
            'codigo_venta': id_venta,
            'fecha': fecha,
            'valor': valor,
            'medio': medio,
            'referencia': ref,
            'observaciones': obs,
            'pendiente_completar': bool(not fecha or not medio)
        })

out = {
    'proyectos': proyectos,
    'lotes': lotes,
    'ventas': ventas,
    'abonos': abonos
}

with open('server/data/excel_dump.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, ensure_ascii=False, indent=2)

print('Datos extraidos exitosamente a server/data/excel_dump.json')
print(f"Proyectos: {len(proyectos)}, Lotes: {len(lotes)}, Ventas: {len(ventas)}, Abonos: {len(abonos)}")
